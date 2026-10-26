"""Fast, resumable LMS sync: create the missing sessions and upload the missing attendance.

Instead of opening every group's sessions page and reading it on screen, this
script reads every group's sessions table straight from the LMS (an HTTP fetch
from inside the logged-in browser, ~0.5s per group, several groups at once),
compares it with the exported CSVs and works out up front exactly what is
missing:

  * create  - a session in "CSV Titles" that does not exist on the LMS yet
  * edit    - a session on the LMS with the wrong title/time (same as edit.py)
  * upload  - a session with an attendance CSV whose attendance is not taken
  * reupload- attendance already taken but the CSV changed since
              (pending_reuploads.json) or FORCE_REUPLOAD is on

The full list is written to exports/lms_missing_sessions.csv before anything
is changed. Then only those items are done, using the same form-filling and
upload code as edit.py / upload.py. Groups with nothing missing cost one fetch.

Progress is saved after every step to exports/lms_fast_sync_state.json. When
a run stops (crash, closed window, Ctrl+C) the next run prints the group,
step and session it stopped at, and continues from there: the fresh scan sees
what was already created/uploaded, so nothing is done twice.

It runs in its own Chrome profile (<profile>_fast) so it never fights with the
edit.py / upload.py profile, and it does not change either of those scripts.

Usage (from the project root):
  python -X utf8 -u lms-bot/fast_sync.py            # scan, report, then sync
  python -X utf8 -u lms-bot/fast_sync.py --plan     # scan and report only
  python -X utf8 -u lms-bot/fast_sync.py --track technical --round 5 --groups CAI5_AIS4_S7,AST5_DAT1_S1

The usual .env / environment settings still apply (LMS_TRACK, ONLY_GROUPS,
LMS_ROUND, DRY_RUN, FORCE_REUPLOAD, START_FROM_DATE).
"""

import argparse
import base64
import csv
import html
import json
import os
import re
import sys
import time
from datetime import date, datetime


def _apply_cli_args() -> argparse.Namespace:
    # edit.py / upload.py read their settings from the environment when they
    # are imported, so the command-line options are turned into env vars first.
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--plan", action="store_true", help="scan and report only, change nothing")
    parser.add_argument("--track", help="technical | nontechnical | all")
    parser.add_argument("--groups", help="comma-separated group names")
    parser.add_argument("--round", help="round number, e.g. 5")
    parser.add_argument(
        "--step",
        choices=["all", "sessions", "upload"],
        default=os.getenv("FAST_SYNC_STEP", "all"),
        help="sessions = create/fix sessions only, upload = upload attendance only, all = both",
    )
    args = parser.parse_args()
    if args.track:
        os.environ["LMS_TRACK"] = args.track
    if args.groups:
        os.environ["ONLY_GROUPS"] = args.groups
    if args.round:
        os.environ["LMS_ROUND"] = args.round
    os.environ.setdefault("DEBUG", "0")
    return args


ARGS = _apply_cli_args()

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from playwright.sync_api import sync_playwright  # noqa: E402

import edit as E  # noqa: E402
import lms_session  # noqa: E402
import upload as U  # noqa: E402
import progress_events as P

PLAN_ONLY = ARGS.plan or E.DRY_RUN or U.DRY_RUN
# Set by Run Everything: the export is stage 1, this script runs stages 2-4.
FULL_RUN_STAGES = os.getenv("FULL_RUN_STAGES", "").strip() == "1"


def print_stage(number, title, profile_name):
    if FULL_RUN_STAGES:
        print("\n" + "=" * 70)
        print(f"[FULL RUN] Stage {number}/4: {title} ({profile_name})")
        print("=" * 70)


DO_SESSIONS = ARGS.step in ("all", "sessions")
DO_UPLOADS = ARGS.step in ("all", "upload")
SCAN_CONCURRENCY = max(1, int(os.getenv("FAST_SCAN_CONCURRENCY", "4")))
PAUSE_MS = int(os.getenv("FAST_PAUSE_MS", "300"))
# FAST_SYNC_HTTP=0 turns the direct form posts off (everything on screen).
USE_HTTP = os.getenv("FAST_SYNC_HTTP", "1").strip().lower() not in {"0", "false", "no"}

EXPORTS_ROOT = U.EXPORTS_ROOT
STATE_PATH = os.path.join(EXPORTS_ROOT, "lms_fast_sync_state.json")
PLAN_CSV_PATH = os.path.join(EXPORTS_ROOT, "lms_missing_sessions.csv")


class BlockedError(RuntimeError):
    """The LMS firewall answered with its "page blocked" screen."""


# The account the browser is signed in with, so a lost session can be renewed.
CURRENT_ACCOUNT = {"email": "", "password": ""}

# Errors that only mean the request was cut off (the page reloaded or the
# connection dropped for a moment): reading again is safe.
_TRANSIENT = (
    "Execution context was destroyed",
    "Failed to fetch",
    "most likely because of a navigation",
    "Cannot find context with specified id",
    "NetworkError",
    "net::ERR_",
    # The LMS logged the account out (e.g. a saved login that expired):
    # recover_page signs in again, then the request is sent again.
    "LMS session expired",
    # A request the LMS did not answer in time (fetch timeout).
    "TimeoutError",
    "signal timed out",
)


def is_transient(error) -> bool:
    return any(s in str(error) for s in _TRANSIENT)


def recover_page(page, attempt):
    """Waits a little, lets the page settle and signs in again if needed."""
    time.sleep(2 * attempt)
    try:
        page.wait_for_load_state("domcontentloaded", timeout=15000)
    except Exception:
        pass
    try:
        # Load a fresh page: a page left open can still show a session the
        # LMS has already ended.
        page.goto(f"{lms_session.LMS_BASE_URL}/my/")
        E.wait_dom_ready(page)
        if "login/index.php" in page.url.lower() or not lms_session.get_sesskey(page):
            print("🔐 The LMS login expired; signing in again...")
            E.ensure_logged_in(page, CURRENT_ACCOUNT["email"], CURRENT_ACCOUNT["password"])
    except Exception as e:
        print(f"⚠️ Could not settle the page yet: {e}")


def evaluate_with_retry(page, js, arg, what, retries=3, result_error=lambda res: None):
    """page.evaluate that survives a page reload / dropped connection.

    Only for requests that are safe to send again (reads, uploads, edits).
    result_error(res) returns the error text inside a result, if any.
    """
    last = None
    for attempt in range(retries + 1):
        if attempt:
            print(f"⚠️ {what}: the request was cut off ({last}); trying again ({attempt}/{retries})...")
            recover_page(page, attempt)
        try:
            res = page.evaluate(js, arg)
        except Exception as e:
            if not is_transient(e):
                raise
            last = e
            continue
        err = result_error(res)
        if err and is_transient(err):
            last = err
            continue
        return res
    raise RuntimeError(f"{what}: {last}")


# =========================
# STATE / CHECKPOINT
# =========================
class State:
    """exports/lms_fast_sync_state.json, rewritten after every step."""

    def __init__(self):
        self.previous = self._load()
        self.data = {
            "startedAt": _now(),
            "updatedAt": _now(),
            "finished": False,
            "planOnly": PLAN_ONLY,
            "track": lms_session.track_label(E.LMS_TRACK),
            "current": None,
            "plan": {},
            "done": [],
            "failed": [],
        }

    @staticmethod
    def _load():
        try:
            with open(STATE_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None

    def report_previous_stop(self):
        prev = self.previous
        if not prev or prev.get("finished") or prev.get("planOnly"):
            return
        cur = prev.get("current") or {}
        print("\n" + "!" * 100)
        print(f"⚠️ The last run (started {prev.get('startedAt')}) did not finish.")
        if cur:
            print(
                f"   It stopped at group {cur.get('group')} | step: {cur.get('step')} | "
                f"session: {cur.get('session')} | at {cur.get('since')}"
            )
        print(
            f"   Done then: {len(prev.get('done') or [])} step(s), failed: {len(prev.get('failed') or [])}. "
            "Continuing: the scan below skips everything already on the LMS."
        )
        print("!" * 100)

    def save(self):
        self.data["updatedAt"] = _now()
        try:
            os.makedirs(EXPORTS_ROOT, exist_ok=True)
            tmp = STATE_PATH + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(self.data, f, ensure_ascii=False, indent=2)
            # Windows can hold the file for a moment (antivirus, a viewer):
            # try again briefly, then write it in place.
            for _ in range(5):
                try:
                    os.replace(tmp, STATE_PATH)
                    return
                except PermissionError:
                    time.sleep(0.2)
            with open(STATE_PATH, "w", encoding="utf-8") as f:
                json.dump(self.data, f, ensure_ascii=False, indent=2)
            os.remove(tmp)
        except Exception as e:
            print(f"⚠️ Could not save {os.path.basename(STATE_PATH)}: {e}")

    def at(self, group, step, session=""):
        self.data["current"] = {"group": group, "step": step, "session": session, "since": _now()}
        self.save()

    def done(self, group, step, session, item=None):
        self.data["done"].append(
            {"group": group, "step": step, "session": session, "at": _now(), **_item_details(item)}
        )
        self.save()

    def failed(self, group, step, session, reason, item=None):
        self.data["failed"].append({
            "group": group, "step": step, "session": session, "reason": str(reason)[:300],
            "at": _now(), **_item_details(item),
        })
        self.save()

    def finish(self):
        self.data["finished"] = True
        self.data["current"] = None
        self.save()


def _item_details(item):
    """What the Excel report shows for one session (see write_results_excel)."""
    if not item:
        return {}
    d = item.get("date")
    return {
        "profile": (E._PROFILE or {}).get("name", ""),
        "date": d.isoformat() if d else "",
        "time": f"{item['start']}-{item['end']}" if item.get("start") else "",
        "title": str(item.get("title", "")),
        "csv_file": os.path.basename(item.get("csv_path", "") or ""),
    }


def _now():
    return datetime.now().isoformat(timespec="seconds")


# =========================
# FAST TABLE SCAN
# =========================
# Same row reading as lms_session.parse_sessions_table, but on HTML fetched in
# the background (no page navigation). Cell texts are joined with spaces,
# which gives the same text as innerText once spaces are collapsed.
_SCAN_JS = """async ([urls, concurrency]) => {
  const out = new Array(urls.length);
  let next = 0;
  const textOf = (el) => {
    const parts = [];
    const walker = document.createTreeWalker(el, NodeFilter.SHOW_TEXT);
    let node;
    while ((node = walker.nextNode())) parts.push(node.nodeValue);
    return parts.join(' ');
  };
  const abs = (a) => {
    const h = a.getAttribute('href') || '';
    try { return h ? new URL(h, location.href).href : ''; } catch (e) { return h; }
  };
  const label = (a) => `${a.getAttribute('aria-label') || ''} ${a.getAttribute('title') || ''}`.toLowerCase();
  async function one(i) {
    try {
      const r = await fetch(urls[i], { credentials: 'include', signal: AbortSignal.timeout(60000) });
      const html = await r.text();
      if (html.includes('has been blocked') && html.includes('Attack ID')) { out[i] = { error: 'blocked' }; return; }
      if (r.url.includes('/login/index.php')) { out[i] = { error: 'login' }; return; }
      const doc = new DOMParser().parseFromString(html, 'text/html');
      const table = doc.querySelector('table.generaltable');
      const rows = table ? Array.from(table.querySelectorAll('tbody tr')) : [];
      out[i] = {
        status: r.status,
        has_table: !!table,
        rows: rows.map((row, index) => {
          const links = Array.from(row.querySelectorAll('a'));
          const takeLinks = links.filter((a) => (a.getAttribute('href') || '').includes('/mod/attendance/take.php'));
          const changeLink = takeLinks.find((a) =>
            label(a).includes('change attendance') ||
            Array.from(a.querySelectorAll('img')).some((img) => (img.getAttribute('src') || '').includes('redo'))
          );
          const takeLink = takeLinks.find((a) =>
            label(a).includes('take attendance') || a.querySelector("i[class*='fa-play']")
          );
          const editLink = links.find((a) => label(a).includes('edit session'));
          return {
            index,
            text: Array.from(row.cells).map(textOf).join(' ').replace(/\\s+/g, ' ').trim(),
            take_href: takeLink ? abs(takeLink) : '',
            change_href: changeLink ? abs(changeLink) : '',
            edit_href: editLink ? abs(editLink) : '',
          };
        }),
      };
    } catch (e) {
      out[i] = { error: String(e) };
    }
  }
  async function worker() { while (next < urls.length) { const i = next++; await one(i); } }
  await Promise.all(Array.from({ length: Math.min(concurrency, urls.length) }, worker));
  return out;
}"""


def scan_tables(page, urls):
    """Returns one list of rows per URL (same keys as parse_sessions_table)."""
    if not urls:
        return []
    results = evaluate_with_retry(
        page, _SCAN_JS, [list(urls), SCAN_CONCURRENCY], "reading the sessions tables",
        result_error=lambda res: next(
            ("the LMS session expired (login page)" if r.get("error") == "login"
             else f"NetworkError: the LMS server answered HTTP {r.get('status')}" if (r.get("status") or 0) >= 500
             else r.get("error")
             for r in res or []
             if r and (r.get("error") == "login" or is_transient(r.get("error", "")) or (r.get("status") or 0) >= 500)),
            None,
        ),
    )
    tables = []
    for url, res in zip(urls, results):
        err = (res or {}).get("error")
        if err == "blocked":
            raise BlockedError(
                "The LMS firewall blocked a request (\"Web Page Blocked\"). Stopped; wait a "
                "while before running again."
            )
        if err == "login":
            raise RuntimeError("The LMS session expired (redirected to the login page).")
        if err:
            raise RuntimeError(f"Could not read {url}: {err}")
        if not res.get("has_table") and res.get("status") != 200:
            raise RuntimeError(f"Could not read {url}: HTTP {res.get('status')}")
        rows = []
        for row in res.get("rows") or []:
            hrefs = " ".join([row.get("edit_href", ""), row.get("change_href", ""), row.get("take_href", "")])
            match = re.search(r"sessionid=(\d+)", hrefs)
            row["session_id"] = match.group(1) if match else ""
            row["attendance_taken"] = bool(row.get("change_href"))
            rows.append(row)
        tables.append(rows)
    return tables


def lms_sessions_from_rows(rows):
    """Same output as edit.build_lms_sessions, from scanned rows."""
    out = []
    for row in rows:
        txt = " ".join(row["text"].split())
        lowered = txt.lower()
        if not txt or "no sessions" in lowered or "nothing to display" in lowered or "no data available" in lowered:
            continue
        d = E.extract_date_from_row_text(txt)
        if not d:
            continue
        title = E.extract_title_from_row_text(txt)
        out.append({
            "idx_1based": row["index"] + 1,
            "date_obj": d,
            "title": title,
            "title_norm": E.normalize_text(title),
            "row_text": txt[:500],
            "attendance_taken": row["attendance_taken"],
            "session_id": row["session_id"],
            "edit_href": row["edit_href"],
            "key": row["session_id"] or f"row:{row['index']}:{txt[:120]}",
        })
    return out


# =========================
# GROUPS
# =========================
def edit_profile_for(track):
    return next(p for p in E.PROFILES if p["track"] == track)


def upload_profile_for(track):
    return next(p for p in U.PROFILES if p["track"] == track)


def use_profiles(track):
    E._PROFILE = edit_profile_for(track)
    U._PROFILE = upload_profile_for(track)


_GROUP_LIST_JS = """async (urls) => Promise.all(urls.map(async (url) => {
  try {
    const r = await fetch(url, { credentials: 'include', signal: AbortSignal.timeout(60000) });
    const html = await r.text();
    if (html.includes('has been blocked') && html.includes('Attack ID')) return { blocked: true };
    const doc = new DOMParser().parseFromString(html, 'text/html');
    const out = {};
    for (const o of doc.querySelectorAll('select[name=group] option')) {
      const m = (o.textContent || '').trim().match(/^Group:\\s*(\\S+)$/i);
      if (m && Number(o.value) > 0) out[m[1]] = Number(o.value);
    }
    return { groups: out };
  } catch (e) {
    return { groups: {} };
  }
}))"""


def course_module_ids(track):
    if E.COURSE_MODULE_ID_OVERRIDE:
        return [str(E.COURSE_MODULE_ID_OVERRIDE)]
    ep = edit_profile_for(track)
    return [str(ep["COURSE_MODULE_ID"]), str(ep["ROUND_5_COURSE_MODULE_ID"])]


def lms_group_numbers(page, track):
    """{course module id: {group name: LMS group number}} from the LMS group list."""
    cmids = course_module_ids(track)
    urls = [f"{lms_session.LMS_BASE_URL}/mod/attendance/manage.php?id={c}&view={E.VIEW}" for c in cmids]
    results = evaluate_with_retry(page, _GROUP_LIST_JS, urls, "reading the LMS group list")
    if any(r.get("blocked") for r in results):
        raise BlockedError(
            "The LMS firewall blocked a request (\"Web Page Blocked\"). Stopped; wait a "
            "while before running again."
        )
    return {c: (r.get("groups") or {}) for c, r in zip(cmids, results)}


def group_names_for_track(track):
    """Groups that have a titles CSV and/or an attendance folder."""
    ep, up = edit_profile_for(track), upload_profile_for(track)
    names = set(up["GROUP_TO_NUM"]) | set(ep["GROUP_TO_NUM"])
    titles_dir = ep["CSV_DIR"]
    if os.path.isdir(titles_dir):
        for fname in os.listdir(titles_dir):
            if fname.lower().endswith(".csv"):
                names.add(E.normalize_group_name_from_filename(fname))
    base = up["EXPORTS_BASE_DIR"]
    if os.path.isdir(base):
        for fname in os.listdir(base):
            if fname != os.path.basename(titles_dir) and os.path.isdir(os.path.join(base, fname)):
                names.add(fname)

    if E.LMS_ROUND:
        names = {n for n in names if re.match(r"^[A-Za-z]{3,4}" + re.escape(E.LMS_ROUND), n)}
    if E.ONLY_GROUPS:
        names = {n for n in names if n in E.ONLY_GROUPS}
        names.update(n for n in E.ONLY_GROUPS if n in up["GROUP_TO_NUM"] or n in ep["GROUP_TO_NUM"])
    return sorted(names)


def groups_for_track(track, names, lms_numbers=None):
    """Each group with its LMS group number.

    The number comes from the LMS's own group list when it has the group (so a
    group missing from the bots' lists is never skipped); the lists in
    upload.py / edit.py are the fallback.
    """
    ep, up = edit_profile_for(track), upload_profile_for(track)
    titles_dir = ep["CSV_DIR"]
    lms_numbers = lms_numbers or {}
    mapped, unmapped = [], []
    for name in names:
        listed = up["GROUP_TO_NUM"].get(name) or ep["GROUP_TO_NUM"].get(name)
        cmid = cmid_from_url(U.build_target_url(listed or 0, name))
        on_lms = (lms_numbers.get(cmid) or {}).get(name)
        if on_lms and listed and int(listed) != int(on_lms):
            print(f"ℹ️ {name}: the LMS lists it as group {on_lms} (the bot list says {listed}); using {on_lms}.")
        elif on_lms and not listed:
            print(f"ℹ️ {name}: not in the bot lists; found on the LMS as group {on_lms}.")
        num = on_lms or listed
        if num:
            titles_csv = os.path.join(titles_dir, f"{name}.csv")
            if not os.path.isfile(titles_csv):
                titles_csv = os.path.join(titles_dir, f"{name}_sessions.csv")
            mapped.append({
                "name": name,
                "num": num,
                "titles_csv": titles_csv if os.path.isfile(titles_csv) else "",
            })
        else:
            unmapped.append(name)
    return mapped, unmapped


# =========================
# PLANNING
# =========================
def plan_titles(group, lms_sessions):
    """What edit.py would do for this group's titles CSV, decided up front."""
    actions, notes = [], []
    if not group["titles_csv"]:
        return P.select(actions), P.select(notes)
    rows_csv = E.read_csv_rows(group["titles_csv"])
    if not rows_csv:
        return P.select(actions), P.select(notes)
    target_dates = E.build_target_dates_with_rollover(rows_csv)

    items = []
    for idx, item in enumerate(rows_csv, start=1):
        title = item.get("title", "").strip()
        start24, end24 = E.extract_time_range_from_title(title)
        base = {"date": target_dates[idx - 1], "start": start24, "end": end24, "title": title, "csv_row": idx}
        if not start24 or not end24:
            notes.append({**base, "action": "problem", "reason": "the title in CSV Titles has no time range (e.g. 6:00 pm - 9:00 pm); fix the title and run again"})
            continue
        items.append(base)

    # Pass 1: every CSV row that is already exactly on the LMS claims its
    # session first, so pass 2 can never move a session another row needs.
    claimed, rest, seen = set(), [], set()
    for base in items:
        key = (base["date"], base["start"], base["end"], E.normalize_text(base["title"]))
        exact = next(
            (s for s in lms_sessions
             if s["key"] not in claimed
             and E.is_same_session_exact(s, base["date"], base["title"], base["start"], base["end"])),
            None,
        )
        if exact:
            claimed.add(exact["key"])
        elif key in seen:
            pass  # the same row twice in CSV Titles: one session is enough
        else:
            rest.append(base)
        seen.add(key)

    # Pass 2: fix a session on the same day (same time first), or create it.
    for base in rest:
        candidate, how = E.choose_lms_session(lms_sessions, base["date"], base["start"], base["end"], claimed)
        if how == "conflict":
            # Two CSV rows want the same date and time with different titles.
            notes.append({**base, "action": "problem", "reason": "two rows in CSV Titles have this date and time with different titles; fix CSV Titles"})
            continue
        if candidate:
            claimed.add(candidate["key"])
            reason = ""
            if candidate.get("attendance_taken"):
                # Same session (nothing else in CSV Titles claims it): only its
                # title/time are corrected, the attendance on it stays.
                reason = "attendance already taken; only the title/time are corrected, the attendance stays"
            actions.append({**base, "action": "edit", "candidate": candidate, "reason": reason})
        else:
            actions.append({**base, "action": "create"})
    return P.select(actions), P.select(notes)


def plan_uploads(group, rows, planned_csvs=()):
    """What upload.py would do for this group, decided up front.

    Every attendance CSV in the group's folder must end up on an LMS session:
    one that is not matched by any LMS row (or by a session about to be
    created, planned_csvs) is reported instead of being silently left out.
    """
    actions, notes = [], []
    covered = {os.path.normcase(os.path.abspath(p)) for p in planned_csvs if p}
    csv_start, csv_end = U.get_csv_date_range_for_group(group["name"])
    if csv_start is None and not U.GLOBAL_START_FROM_DATE:
        return P.select(actions), P.select(notes)
    start_from = U.GLOBAL_START_FROM_DATE or csv_start

    for row in rows:
        text = row["text"]
        if not text:
            continue
        row_date = U.extract_date_from_row_text(text)
        if not row_date:
            if not row["attendance_taken"]:
                notes.append({"date": None, "action": "problem", "title": text[:120], "reason": "could not read the date of this LMS row"})
            continue
        if start_from and row_date < start_from:
            continue
        if csv_end and row_date > csv_end:
            continue
        taken = row["attendance_taken"]
        row_group = U.extract_group_from_row_text(text) or group["name"]
        session_title = U.extract_session_title_from_row_text(text)
        session_type = U.extract_session_type_from_title(session_title)
        base = {"date": row_date, "title": session_title, "row_group": row_group, "session_id": row["session_id"]}
        try:
            csv_path = U.find_csv_for_session(row_group, row_date, session_type)
        except RuntimeError as e:
            if not taken:
                notes.append({**base, "action": "problem", "reason": f"no attendance CSV for this LMS session: {e}"})
            continue
        covered.add(os.path.normcase(os.path.abspath(csv_path)))
        if not taken and not row["take_href"]:
            notes.append({**base, "action": "problem", "csv_path": csv_path, "reason": "the LMS row has no Take attendance link"})
            continue
        key = U.pending_reupload_key(csv_path)
        if taken:
            if key in U.PENDING_REUPLOADS or U.FORCE_REUPLOAD:
                reason = "CSV changed since the last upload" if key in U.PENDING_REUPLOADS else "re-upload option is on"
                actions.append({**base, "action": "reupload", "csv_path": csv_path, "href": row["change_href"], "reason": reason})
            continue
        actions.append({**base, "action": "upload", "csv_path": csv_path, "href": row["take_href"]})

    group_dir = os.path.join(U._p("EXPORTS_BASE_DIR"), group["name"])
    date_re = re.compile(r"^(\d{4})_(\d{2})_(\d{2})")
    for fname in sorted(os.listdir(group_dir)) if os.path.isdir(group_dir) else []:
        m = date_re.match(fname)
        if not fname.lower().endswith(".csv") or not m:
            continue
        d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        if start_from and d < start_from:
            continue
        path = os.path.join(group_dir, fname)
        if os.path.normcase(os.path.abspath(path)) in covered:
            continue
        notes.append({
            "date": d, "action": "problem", "title": fname, "csv_path": path,
            "reason": "this attendance CSV has no session on the LMS"
            + (" yet: it is created by Sync / Edit Sessions" if group.get("titles_csv") else " and no row in CSV Titles"),
        })
    return P.select(actions), P.select(notes)


def describe(item):
    d = item.get("date")
    when = d.isoformat() if d else "?"
    if item.get("start"):
        when += f" {item['start']}-{item['end']}"
    return f"{when} | {str(item.get('title', ''))[:70]}"


_ACTION_LABELS = {
    "create": "Created",
    "edit": "Edited",
    "upload": "Uploaded",
    "reupload": "Re-uploaded",
}


def write_results_excel(state):
    """exports/uploaded_sessions_<dates>_<time>.xlsx: what this run did.

    Summary (totals per track), Done (every session created / edited /
    uploaded, verified on the LMS) and Failed. Next to the dashboard's
    skipped_sessions report, which lists what could not be exported.
    """
    done, failed = state.data.get("done") or [], state.data.get("failed") or []
    if not done and not failed:
        print("📊 Nothing was created or uploaded in this run, so no uploaded-sessions report.")
        return None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("⚠️ openpyxl is not installed (pip install openpyxl): no Excel report.")
        return None

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = bold, head_fill
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"

    wb = Workbook()

    # ---- Summary: totals per track ----
    profiles = sorted({r.get("profile") or "" for r in done + failed}) or [""]
    summary = []
    for prof in profiles:
        rows = [r for r in done if (r.get("profile") or "") == prof]
        counts = {k: sum(1 for r in rows if r["step"] == k) for k in _ACTION_LABELS}
        summary.append([
            prof or "-", counts["create"], counts["edit"], counts["upload"], counts["reupload"],
            counts["upload"] + counts["reupload"],
            sum(1 for r in failed if (r.get("profile") or "") == prof),
        ])
    summary.append([
        "TOTAL", *[sum(r[i] for r in summary) for i in range(1, 7)],
    ])
    ws = wb.active
    ws.title = "Summary"
    sheet(ws, ["Track", "Created", "Edited", "Uploaded", "Re-uploaded", "Total uploaded", "Failed"], summary,
          [16, 10, 10, 10, 13, 15, 10])
    for cell in ws[ws.max_row]:
        cell.font = bold
    ws.append([])
    ws.append(["Run started", state.data.get("startedAt", "")])
    ws.append(["Run ended", _now()])

    # ---- Done / Failed: one row per session ----
    sheet(
        wb.create_sheet("Done"),
        ["#", "Track", "Group", "Date", "Time", "Action", "Session title", "CSV file", "Done at"],
        [[i, r.get("profile", ""), r["group"], r.get("date", ""), r.get("time", ""),
          _ACTION_LABELS.get(r["step"], r["step"]), r.get("title") or r.get("session", ""),
          r.get("csv_file", ""), r["at"].replace("T", " ")]
         for i, r in enumerate(sorted(done, key=lambda r: (r.get("profile", ""), r["group"], r.get("date", ""))), start=1)],
        [5, 15, 16, 12, 12, 12, 60, 42, 20],
    )
    sheet(
        wb.create_sheet("Failed"),
        ["#", "Track", "Group", "Date", "Time", "Action", "Session title", "CSV file", "Reason", "At"],
        [[i, r.get("profile", ""), r["group"], r.get("date", ""), r.get("time", ""),
          str(r["step"]).capitalize(), r.get("title") or r.get("session", ""),
          r.get("csv_file", ""), r.get("reason", ""), r["at"].replace("T", " ")]
         for i, r in enumerate(failed, start=1)],
        [5, 15, 16, 12, 12, 12, 60, 42, 60, 20],
    )

    start, end = os.getenv("START_FROM_DATE", "").strip(), os.getenv("END_AT_DATE", "").strip()
    dates = f"{start.replace('-', '_')}_to_{end.replace('-', '_')}_" if start and end else ""
    path = os.path.join(EXPORTS_ROOT, f"uploaded_sessions_{dates}{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")
    try:
        wb.save(path)
    except Exception as e:
        print(f"⚠️ Could not save the uploaded-sessions report (is it open in Excel?): {e}")
        return None
    t = summary[-1]
    print(
        f"📊 Uploaded-sessions report: {path}\n"
        f"   TOTAL: created={t[1]} | edited={t[2]} | uploaded={t[3]} | re-uploaded={t[4]} | "
        f"total uploaded={t[5]} | failed={t[6]}"
    )
    return path


def write_plan_csv(plan_rows):
    os.makedirs(EXPORTS_ROOT, exist_ok=True)
    fields = ["profile", "group", "action", "date", "time", "title_or_type", "csv_file", "reason"]
    try:
        with open(PLAN_CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            writer.writeheader()
            writer.writerows(plan_rows)
        print(f"📋 Missing list saved: {PLAN_CSV_PATH} ({len(plan_rows)} rows)")
    except Exception as e:
        print(f"⚠️ Could not save {PLAN_CSV_PATH} (is it open in Excel?): {e}")


def plan_row(profile_name, group_name, item):
    d = item.get("date")
    return {
        "profile": profile_name,
        "group": group_name,
        "action": (
            f"{item['action']} + upload" if item["action"] in ("create", "edit") and item.get("csv_path")
            else f"{item['action']} only" if item["action"] in ("create", "edit", "upload")
            else item["action"]
        ),
        "date": d.isoformat() if d else "",
        "time": f"{item['start']}-{item['end']}" if item.get("start") else "",
        "title_or_type": item.get("title", ""),
        "csv_file": os.path.basename(item.get("csv_path", "") or ""),
        "reason": item.get("reason", ""),
    }


# =========================
# ACTIONS ON THE LMS (browser, same steps as edit.py / upload.py)
# =========================
def cmid_from_url(target_url):
    m = re.search(r"[?&]id=(\d+)", target_url)
    return m.group(1) if m else ""


def open_add_form(page, target_url):
    """Opens the Add session form directly (the table page is not loaded first)."""
    page.goto(f"{lms_session.LMS_BASE_URL}/mod/attendance/sessions.php?id={cmid_from_url(target_url)}&action=1")
    E.wait_dom_ready(page)
    E.dismiss_popups(page)
    try:
        expand_all = page.locator("button, a").filter(has_text=re.compile(r"expand all", re.I)).first
        if expand_all.count() > 0 and expand_all.is_visible(timeout=1500):
            expand_all.click()
            page.wait_for_timeout(400)
    except Exception:
        pass
    try:
        page.locator("select[name='groups[]'], #id_groups").first.wait_for(state="visible", timeout=8000)
    except Exception:
        # Direct link refused: fall back to the Add session button on the table.
        E.goto_add_session_page(page, target_url)


def do_create(page, group, target_url, item):
    open_add_form(page, target_url)
    E.set_group_for_add_session(page, group["name"])
    E.set_date_dropdowns(page, item["date"])
    E.set_time_dropdowns(page, item["start"], item["end"])
    E.set_description_tinymce(page, item["title"])
    E.submit_add_session(page)


def do_edit(page, target_url, item):
    E.open_edit_form(page, item["candidate"], target_url)
    E.set_description_tinymce(page, item["title"])
    E.set_date_dropdowns(page, item["date"])
    E.set_time_dropdowns(page, item["start"], item["end"])
    E.save_changes(page)


def do_upload(page, item):
    page.goto(item["href"])
    U.wait_dom_ready(page)
    U.dismiss_popups(page)
    U.upload_attendance_csv_for_current_session(page, item["csv_path"], item["row_group"])


# =========================
# DIRECT FORM SUBMISSION (no clicking)
# The same Moodle forms the browser shows are fetched, filled and posted from
# inside the logged-in page. Every step checks the LMS answer before the next
# one (form found, file stored, "Sessions updated" shown), so nothing depends
# on how fast a page renders. Any surprise raises, and the caller falls back
# to the on-screen steps above.
# =========================
_HTTP_COMMON_JS = """
  const parse = (h) => new DOMParser().parseFromString(h, 'text/html');
  const blocked = (h) => h.includes('has been blocked') && h.includes('Attack ID');
  const texts = (doc, sel) => Array.from(doc.querySelectorAll(sel))
    .map((e) => (e.textContent || '').replace(/\\s+/g, ' ').trim()).filter(Boolean);
  const get = async (url) => {
    const r = await fetch(url, { credentials: 'include', signal: AbortSignal.timeout(60000) });
    return { url: r.url, status: r.status, html: await r.text() };
  };
  const post = async (url, body) => {
    const r = await fetch(url, { method: 'POST', body, credentials: 'include', signal: AbortSignal.timeout(120000) });
    return { url: r.url, status: r.status, html: await r.text() };
  };
  const check = (res, what) => {
    if (blocked(res.html)) throw { blocked: true };
    if (res.url.includes('/login/index.php')) throw { error: 'the LMS session expired (login page)' };
    if (res.status >= 400) throw { error: `${what}: HTTP ${res.status}` };
    return parse(res.html);
  };
"""

# Add session (marker addsession, with the group) and Edit session (marker
# updatesession, the group stays as it is) are the same Moodle form.
_HTTP_SESSION_FORM_JS = """async ([formUrl, marker, groupValue, groupName, y, m, d, sh, sm, eh, em, descHtml]) => {
""" + _HTTP_COMMON_JS + """
  try {
    const doc = check(await get(formUrl), 'session form');
    const form = Array.from(doc.querySelectorAll('form'))
      .find((f) => f.querySelector(`[name=${marker}]`));
    if (!form) return { error: 'the session form was not found' };
    const fd = new FormData(form);
    if (groupValue !== null) {
      const select = form.querySelector('select[name="groups[]"]');
      const options = select ? Array.from(select.options) : [];
      const option = options.find((o) => o.value === String(groupValue))
        || options.find((o) => o.textContent.trim() === groupName);
      if (!option) return { error: `group ${groupName} is not in the form's group list` };
      fd.delete('groups[]');
      fd.append('groups[]', option.value);
    }
    const values = {
      'sessiondate[day]': d, 'sessiondate[month]': m, 'sessiondate[year]': y,
      'sestime[starthour]': sh, 'sestime[startminute]': sm,
      'sestime[endhour]': eh, 'sestime[endminute]': em,
      'sdescription[text]': descHtml, 'sdescription[format]': 1,
    };
    for (const [k, v] of Object.entries(values)) fd.set(k, String(v));
    fd.set('submitbutton', form.querySelector('[name=submitbutton]').value);

    const res = await post(form.getAttribute('action'), fd);
    const out = check(res, 'saving the session');
    if (out.querySelector(`[name=${marker}]`)) {
      const errors = texts(out, '.invalid-feedback, .error, .alert-danger');
      return { error: 'the LMS refused the form: ' + (errors.join(' | ') || 'no message') };
    }
    return { ok: true };
  } catch (e) {
    if (e && e.blocked) return { blocked: true };
    return { error: (e && e.error) || String(e) };
  }
}"""

_HTTP_UPLOAD_JS = """async ([takeUrl, fileName, fileB64]) => {
""" + _HTTP_COMMON_JS + """
  try {
    // 1) "Upload attendance by CSV" button on the take page.
    let doc = check(await get(takeUrl), 'take attendance page');
    let form = Array.from(doc.querySelectorAll('form'))
      .find((f) => (f.getAttribute('action') || '').includes('/import/marksessions.php'));
    if (!form) return { error: 'the "Upload attendance by CSV" button was not found' };
    let res = await post(form.getAttribute('action'), new FormData(form));
    doc = check(res, 'upload form');
    form = Array.from(doc.querySelectorAll('form'))
      .find((f) => f.querySelector('[name=_qf__mod_attendance_form_import_marksessions]'));
    const fileField = form && form.querySelector('[name=attendancefile]');
    const clientId = (res.html.match(/"client_id":"([^"]+)"/) || [])[1];
    const ctxId = (res.html.match(/ctx_id=(\\d+)/) || [])[1];
    const repoId = (res.html.match(/"(\\d+)":\\{"id":"\\d+","name":"[^"]*","type":"upload"/) || [])[1];
    if (!form || !fileField || !clientId || !ctxId || !repoId) return { error: 'the CSV upload form looks different than expected' };

    // 2) Store the CSV in the form's draft area (what the file picker does).
    const bytes = Uint8Array.from(atob(fileB64), (c) => c.charCodeAt(0));
    const up = new FormData();
    up.append('repo_upload_file', new Blob([bytes], { type: 'text/csv' }), fileName);
    const fields = {
      title: fileName, author: '', license: 'unknown', itemid: fileField.value, repo_id: repoId,
      p: '', page: '', env: 'filepicker', sesskey: form.querySelector('[name=sesskey]').value,
      client_id: clientId, maxbytes: '-1', areamaxbytes: '-1', ctx_id: ctxId, savepath: '/',
    };
    for (const [k, v] of Object.entries(fields)) up.append(k, v);
    res = await post('/repository/repository_ajax.php?action=upload', up);
    if (blocked(res.html)) return { blocked: true };
    let stored = null;
    try { stored = JSON.parse(res.html); } catch (e) { /* handled below */ }
    if (!stored || stored.error || !stored.id) {
      return { error: 'the CSV was not accepted: ' + ((stored && stored.error) || res.html.slice(0, 200)) };
    }

    // 3) Submit the file -> column mapping page.
    let fd = new FormData(form);
    fd.set('submitbutton', form.querySelector('[name=submitbutton]').value);
    doc = check(await post(form.getAttribute('action'), fd), 'submitting the CSV');
    form = Array.from(doc.querySelectorAll('form'))
      .find((f) => f.querySelector('[name=_qf__mod_attendance_form_import_marksessions_confirm]'));
    const userTo = form && form.querySelector('select[name=userto]');
    if (!form || !userTo || !Array.from(userTo.options).some((o) => o.value === 'email')) {
      const errors = texts(doc, '.invalid-feedback, .error, .alert');
      return { error: 'the column mapping page did not open: ' + (errors.join(' | ') || 'no message') };
    }

    // 4) Moodle user field = Email address, then confirm: this saves the attendance.
    fd = new FormData(form);
    fd.set('userto', 'email');
    fd.set('submitbutton', form.querySelector('[name=submitbutton]').value);
    doc = check(await post(form.getAttribute('action'), fd), 'confirming the import');
    const alerts = texts(doc, '.alert, [role=alert]');
    if (doc.querySelector('[name=_qf__mod_attendance_form_import_marksessions_confirm]')) {
      return { error: 'the import was not confirmed: ' + (alerts.join(' | ') || 'no message') };
    }
    if (!alerts.some((a) => /sessions? updated/i.test(a))) {
      return { error: 'no "Sessions updated" message after the import: ' + (alerts.join(' | ') || 'no message') };
    }
    return { ok: true, alerts };
  } catch (e) {
    if (e && e.blocked) return { blocked: true };
    return { error: (e && e.error) || String(e) };
  }
}"""


class _AlertTexts:
    """Lets upload.collect_missing_from_wavz read alerts that were fetched, not shown."""

    def __init__(self, texts):
        self.texts = list(texts or [])

    def locator(self, _selector):
        return self

    def all_inner_texts(self):
        return self.texts


def _http_result(res, what):
    if res.get("blocked"):
        raise BlockedError(
            "The LMS firewall blocked a request (\"Web Page Blocked\"). Stopped; wait a "
            "while before running again."
        )
    if not res.get("ok"):
        raise RuntimeError(f"{what}: {res.get('error') or 'unknown error'}")
    return res


def _http_session_form(page, form_url, marker, group_value, group_name, item, what):
    sh, sm = map(int, item["start"].split(":"))
    eh, em = map(int, item["end"].split(":"))
    d = item["date"]
    desc = "<p>" + html.escape(item["title"], quote=False) + "</p>"
    args = [form_url, marker, group_value, group_name, d.year, d.month, d.day, sh, sm, eh, em, desc]
    if group_value is None:
        # An edit sets the same values again, so it is safe to repeat.
        res = evaluate_with_retry(page, _HTTP_SESSION_FORM_JS, args, what, result_error=lambda r: r.get("error"))
    else:
        # A create is not sent twice: the check after it tells whether it was saved.
        res = page.evaluate(_HTTP_SESSION_FORM_JS, args)
    _http_result(res, what)


def http_create(page, group, target_url, item):
    # The LMS group number is what the form's group list uses.
    add_url = f"{lms_session.LMS_BASE_URL}/mod/attendance/sessions.php?id={cmid_from_url(target_url)}&action=1"
    _http_session_form(
        page, add_url, "_qf__mod_attendance_form_addsession",
        str(group["num"]), group["name"], item, "direct create",
    )


def http_edit(page, item):
    href = item["candidate"].get("edit_href")
    if not href:
        raise RuntimeError("direct edit: the LMS row has no Edit link")
    _http_session_form(
        page, href, "_qf__mod_attendance_form_updatesession", None, "", item, "direct edit",
    )


def http_upload(page, item):
    with open(item["csv_path"], "rb") as f:
        data = base64.b64encode(f.read()).decode("ascii")
    res = evaluate_with_retry(
        page, _HTTP_UPLOAD_JS, [item["href"], os.path.basename(item["csv_path"]), data],
        "direct upload", result_error=lambda r: r.get("error"),
    )
    _http_result(res, "direct upload")
    U.collect_missing_from_wavz(_AlertTexts(res.get("alerts")), item["row_group"], item["csv_path"])


# =========================
# ONE GROUP
# =========================
def _run_session_item(page, group, target_url, item, direct):
    if item["action"] == "create":
        print(f"➕ Creating{'' if direct else ' (on screen)'}: {describe(item)}")
        if direct:
            http_create(page, group, target_url, item)
        else:
            do_create(page, group, target_url, item)
    else:
        note = " (attendance on it stays)" if item.get("candidate", {}).get("attendance_taken") else ""
        print(f"✏️ Fixing{'' if direct else ' (on screen)'}{note}: {describe(item)}")
        if direct:
            http_edit(page, item)
        else:
            do_edit(page, target_url, item)


def _run_upload_item(page, item, direct):
    verb = "♻️ Re-uploading" if item["action"] == "reupload" else "⬆️ Uploading"
    print(f"{verb}{'' if direct else ' (on screen)'}: {describe(item)} <- {os.path.basename(item['csv_path'])}")
    if direct:
        http_upload(page, item)
    else:
        do_upload(page, item)


def sync_sessions(page, state, group, target_url, title_actions, stats):
    """Creates / fixes the sessions, verifies them on the LMS and tries every
    miss once more on screen. Returns the fresh table rows."""
    name = group["name"]
    pending = list(title_actions)
    rows = None
    for round_no in (1, 2):
        direct = USE_HTTP and round_no == 1
        for item in pending:
            state.at(name, item["action"], describe(item))
            P.session(name, item, "sync", "running")
            t = time.time()
            try:
                _run_session_item(page, group, target_url, item, direct)
                print(f"   saved in {time.time() - t:.1f}s")
            except BlockedError:
                raise
            except Exception as e:
                item["error"] = str(e)
                print(f"⚠️ {e}")
            page.wait_for_timeout(PAUSE_MS)

        # One fresh read verifies every change (and gives the new sessions'
        # links for the uploads).
        rows = scan_tables(page, [target_url])[0]
        fresh = lms_sessions_from_rows(rows)
        claimed = set()
        missing = []
        for item in pending:
            found = next(
                (s for s in fresh if s["key"] not in claimed
                 and E.is_same_session_exact(s, item["date"], item["title"], item["start"], item["end"])),
                None,
            )
            if found:
                claimed.add(found["key"])
                stats["created" if item["action"] == "create" else "edited"] += 1
                state.done(name, item["action"], describe(item), item=item)
                P.session(name, item, "sync" if item["action"] in ("create", "edit") else "upload", "done")
            else:
                missing.append(item)
        if not missing:
            break
        if round_no == 1:
            print(f"🔁 {len(missing)} session(s) not right on the LMS yet; trying them again on screen.")
            for item in missing:
                if item["action"] == "create":
                    # If the first try did save something at that date and
                    # time, fix that session instead of adding a second one.
                    candidate, how = E.choose_lms_session(fresh, item["date"], item["start"], item["end"], claimed)
                    if how == "time":
                        claimed.add(candidate["key"])
                        item["action"], item["candidate"] = "edit", candidate
            pending = missing
        else:
            for item in missing:
                stats["failed"] += 1
                reason = item.get("error") or "not found on the LMS after saving"
                print(f"❌ Could not {item['action']} {describe(item)}: {reason}")
                state.failed(name, item["action"], describe(item), reason, item=item)
                P.session(name, item, "sync" if item["action"] in ("create", "edit") else "upload", "error", reason)
    return rows


def sync_uploads(page, state, group, target_url, rows, stats):
    """Uploads the missing attendance, verifies it on the LMS and tries every
    miss once more on screen. Returns the problems found while planning."""
    name = group["name"]
    upload_actions, notes = plan_uploads(group, rows)
    pending = list(upload_actions)
    for round_no in (1, 2):
        if not pending:
            break
        direct = USE_HTTP and round_no == 1
        for item in pending:
            state.at(name, item["action"], f"{describe(item)} | {os.path.basename(item['csv_path'])}")
            P.session(name, item, "upload", "running")
            t = time.time()
            try:
                _run_upload_item(page, item, direct)
                print(f"   uploaded in {time.time() - t:.1f}s")
                key = U.pending_reupload_key(item["csv_path"])
                info = U.PENDING_REUPLOADS.pop(key, None)
                if info is not None:
                    item["pending"] = (key, info)
                    U.save_pending_reuploads()
            except BlockedError:
                raise
            except Exception as e:
                # Uploading the same CSV again only rewrites the same marks,
                # so trying again is always safe.
                item["error"] = str(e)
                print(f"⚠️ {e}")
            page.wait_for_timeout(PAUSE_MS)

        after = {r["session_id"]: r for r in scan_tables(page, [target_url])[0] if r["session_id"]}
        missing = []
        for item in pending:
            row = after.get(item["session_id"])
            if row and row["attendance_taken"] and not item.get("error"):
                stats["reuploaded" if item["action"] == "reupload" else "uploaded"] += 1
                state.done(name, item["action"], describe(item), item=item)
                P.session(name, item, "sync" if item["action"] in ("create", "edit") else "upload", "done")
            else:
                if row and row.get("change_href"):
                    item["href"] = row["change_href"]
                missing.append(item)
        if missing and round_no == 1:
            print(f"🔁 {len(missing)} upload(s) not confirmed; trying them again on screen.")
            for item in missing:
                item.pop("error", None)
        elif missing:
            for item in missing:
                stats["failed"] += 1
                reason = item.get("error") or "attendance not shown as taken after the upload"
                print(f"❌ Could not upload {describe(item)}: {reason}")
                state.failed(name, item["action"], describe(item), reason, item=item)
                P.session(name, item, "sync" if item["action"] in ("create", "edit") else "upload", "error", reason)
                key, info = item.get("pending") or (None, None)
                if info is not None:
                    U.PENDING_REUPLOADS[key] = info
                    U.save_pending_reuploads()
        pending = missing
    return notes


def new_stats():
    return {"created": 0, "edited": 0, "uploaded": 0, "reuploaded": 0, "failed": 0}


def add_stats(totals, stats):
    for k in totals:
        totals[k] += stats[k]


def print_group_done(name, stats, started):
    print(
        f"✅ DONE {name} | created={stats['created']} | edited={stats['edited']} | "
        f"uploaded={stats['uploaded']} | reuploaded={stats['reuploaded']} | failed={stats['failed']} "
        f"| {time.time() - started:.1f}s"
    )


# =========================
# MAIN
# =========================
def open_browser(p, email, password):
    # Its own profile folder, so it never shares (or locks) the one used by
    # edit.py / upload.py. The LMS firewall blocks headless browsers, so the
    # window is visible like the other bots.
    user_data_dir = lms_session.user_data_dir_for(email) + "_fast"
    lms_session.harden_chrome_profile(user_data_dir)
    context = p.chromium.launch_persistent_context(
        user_data_dir=user_data_dir,
        headless=False,
        channel="chrome",
        args=lms_session.CHROME_LAUNCH_ARGS,
    )
    page = context.new_page()
    page.set_default_timeout(E.DEFAULT_TIMEOUT_MS)
    E.ensure_logged_in(page, email, password)
    CURRENT_ACCOUNT.update(email=email, password=password)
    return context, page


def main():
    run_started = time.time()
    state = State()
    state.report_previous_stop()

    step_label = {"all": "sessions + upload", "sessions": "sessions only", "upload": "upload only"}[ARGS.step]
    print(
        f"🎛️ Track: {lms_session.track_label(E.LMS_TRACK)} | step: {step_label} | "
        f"mode: {'CHECK ONLY (nothing is changed)' if PLAN_ONLY else 'check, then sync'}"
    )
    if U.FORCE_REUPLOAD:
        print("♻️ Re-upload option is ON: sessions that already have attendance are uploaded again.")
    U.load_student_status()
    U.load_existing_missing_from_wavz()
    U.load_pending_reuploads()

    tracks = [p["track"] for p in E.PROFILES if lms_session.track_is_selected(p["track"])]
    if not tracks:
        print("❌ No profile matches the selected track.")
        return

    plan_rows = []
    problems = []  # (profile, group, note): things that cannot be done automatically
    group_upload_problems = {}  # (profile, group) -> latest upload problems
    totals = {"created": 0, "edited": 0, "uploaded": 0, "reuploaded": 0, "failed": 0}
    stopped = None

    with sync_playwright() as p:
        context = None
        current_email = None
        try:
            for track in tracks:
                use_profiles(track)
                P.TRACK = track
                P.emit("phase", phase="check", status="running")
                profile_name = E._PROFILE["name"]
                names = group_names_for_track(track)
                if not names:
                    print(f"ℹ️ {profile_name}: no groups to check.")
                    continue

                for name in names:
                    P.emit("group", group=name, status="queued", phase="check")
                email, password = lms_session.account_for_track(track)
                if context is not None and email != current_email:
                    context.close()
                    context = None
                if context is None:
                    context, page = open_browser(p, email, password)
                    current_email = email

                groups, unmapped = groups_for_track(track, names, lms_group_numbers(page, track))
                for name in unmapped:
                    print(f"⚠️ {profile_name}: {name} is not a group on the LMS course.")
                    problem = {"action": "problem", "reason": "this group does not exist on the LMS course (ask for it to be added)"}
                    plan_rows.append(plan_row(profile_name, name, problem))
                    problems.append((profile_name, name, problem))
                    P.emit("group", group=name, status="error", reason=problem["reason"])
                if not groups:
                    continue

                # ---- SCAN every group in one go ----
                print_stage(2, "Check missing", profile_name)
                state.at("(scan)", "scan", f"{profile_name}: {len(groups)} groups")
                urls = [U.build_target_url(g["num"], g["name"]) for g in groups]
                t = time.time()
                tables = []
                for offset in range(0, len(groups), SCAN_CONCURRENCY):
                    batch = groups[offset:offset + SCAN_CONCURRENCY]
                    for g in batch:
                        P.emit("group", group=g["name"], status="running", phase="check")
                    tables.extend(scan_tables(page, urls[offset:offset + SCAN_CONCURRENCY]))
                    for g in batch:
                        P.emit("group", group=g["name"], status="scanned", phase="check")
                print(f"\n⚡ {profile_name}: read {len(groups)} group tables in {time.time() - t:.1f}s")

                # ---- PLAN ----
                work = []
                for group, url, rows in zip(groups, urls, tables):
                    P.emit("group", group=group["name"], status="running", phase="check")
                    local_folder = os.path.join(U._p("EXPORTS_BASE_DIR"), group["name"])
                    local_files = [f for f in os.listdir(local_folder) if f.lower().endswith('.csv')] if os.path.isdir(local_folder) else []
                    P.emit("group", group=group["name"], folderCount=len(local_files))
                    for existing in lms_sessions_from_rows(rows):
                        item = {"date": existing["date_obj"], "title": existing["title"]}
                        if P.in_scope(item):
                            P.session(group["name"], item, "check", "done")
                            P.session(group["name"], item, "sync", "done")
                            P.session(group["name"], item, "upload", "done" if existing["attendance_taken"] else "pending")
                    title_actions, title_notes = plan_titles(group, lms_sessions_from_rows(rows))
                    if not DO_SESSIONS:
                        title_actions, title_notes = [], []
                    # One check sorts every session: create/edit only, create/edit
                    # + upload (its attendance CSV is ready and not on the LMS
                    # yet), or upload only.
                    for item in title_actions:
                        if not DO_UPLOADS:
                            continue
                        if item["action"] == "edit" and item["candidate"].get("attendance_taken") and not U.FORCE_REUPLOAD:
                            continue
                        try:
                            session_type = U.extract_session_type_from_title(item["title"])
                            item["csv_path"] = U.find_csv_for_session(group["name"], item["date"], session_type)
                        except Exception:
                            pass
                    upload_actions, upload_notes = [], []
                    if DO_UPLOADS:
                        upload_actions, upload_notes = plan_uploads(
                            group, rows, [i.get("csv_path") for i in title_actions]
                        )
                        # A session being fixed is uploaded with its new title, so it
                        # is counted once, as "edit + upload".
                        edited = {i["candidate"]["session_id"] for i in title_actions if i["action"] == "edit"}
                        upload_actions = [u for u in upload_actions if u["session_id"] not in edited]
                    for item in title_notes:
                        problems.append((profile_name, group["name"], item))
                    group_upload_problems[(profile_name, group["name"])] = upload_notes
                    for item in title_actions + title_notes + upload_actions + upload_notes:
                        plan_rows.append(plan_row(profile_name, group["name"], item))
                    counts = {
                        "create only": sum(1 for i in title_actions if i["action"] == "create" and not i.get("csv_path")),
                        "create + upload": sum(1 for i in title_actions if i["action"] == "create" and i.get("csv_path")),
                        "edit only": sum(1 for i in title_actions if i["action"] == "edit" and not i.get("csv_path")),
                        "edit + upload": sum(1 for i in title_actions if i["action"] == "edit" and i.get("csv_path")),
                        "upload only": sum(1 for i in upload_actions if i["action"] == "upload"),
                        "reupload": sum(1 for i in upload_actions if i["action"] == "reupload"),
                    }
                    for item in title_actions + upload_actions:
                        P.session(group["name"], item, "check", "done")
                        P.session(group["name"], item, "sync" if item["action"] in ("create", "edit") else "upload", "pending")
                    for item in title_notes + upload_notes:
                        P.session(group["name"], item, "check", "error", item.get("reason", ""))
                    P.emit("group", group=group["name"], status="checked", phase="check", planned=counts,
                           reason="No sessions in folder" if not local_files else "")
                    state.data["plan"][group["name"]] = counts
                    if any(counts.values()):
                        work.append((group, url, rows, title_actions, counts))

                print("\n" + "-" * 100)
                print(f"📋 CHECK - {profile_name}: {len(work)} of {len(groups)} groups have something missing")
                for group, _, _, _, counts in work:
                    print(f"   {group['name']:<16} " + " | ".join(f"{k}: {v}" for k, v in counts.items() if v))
                totals_plan = {k: sum(w[4][k] for w in work) for k in (work[0][4] if work else {})}
                if totals_plan:
                    print("   " + "-" * 60)
                    print("   TOTAL            " + " | ".join(f"{k}: {v}" for k, v in totals_plan.items() if v))
                print("-" * 100)
                state.save()

                P.emit("phase", phase="check", status="done")
                if PLAN_ONLY:
                    continue

                # ---- PHASE 1: create / fix the sessions of every group ----
                session_work = [w for w in work if w[3]] if DO_SESSIONS else []
                if DO_SESSIONS:
                    print_stage(3, "Sync sessions (create / edit)", profile_name)
                    if not session_work:
                        print("✅ No session to create or fix.")
                if session_work:
                    P.emit("phase", phase="sync", status="running")
                    print("\n" + "█" * 100)
                    print(f"🗓️ PHASE 1 - {profile_name}: create / fix sessions in {len(session_work)} group(s)")
                    print("█" * 100)
                    for group, url, rows, title_actions, _ in session_work:
                        stats = new_stats()
                        started = time.time()
                        print(f"\n🧩 {group['name']} | {url}")
                        sync_sessions(page, state, group, url, title_actions, stats)
                        print_group_done(group["name"], stats, started)
                        add_stats(totals, stats)

                # ---- PHASE 2: check again, then upload the attendance ----
                if session_work:
                    P.emit("phase", phase="sync", status="done")
                if DO_UPLOADS:
                    print_stage(4, "Upload attendance", profile_name)
                    P.emit("phase", phase="upload", status="running")
                    if session_work:
                        # No second check: only the groups phase 1 changed are read
                        # again, to get the links of the sessions just created.
                        changed = {w[0]["name"] for w in session_work}
                        idx = [i for i, g in enumerate(groups) if g["name"] in changed]
                        t = time.time()
                        fresh_tables = scan_tables(page, [urls[i] for i in idx])
                        tables = list(tables)
                        for i, fresh_rows in zip(idx, fresh_tables):
                            tables[i] = fresh_rows
                        print(f"\n⚡ {profile_name}: read the {len(idx)} changed group(s) again in {time.time() - t:.1f}s")
                    upload_work = []
                    for group, url, rows in zip(groups, urls, tables):
                        actions, notes = plan_uploads(group, rows)
                        group_upload_problems[(profile_name, group["name"])] = notes
                        if actions:
                            upload_work.append((group, url, rows, len(actions)))
                    print("\n" + "█" * 100)
                    print(
                        f"⬆️ PHASE 2 - {profile_name}: upload attendance in {len(upload_work)} group(s), "
                        f"{sum(w[3] for w in upload_work)} session(s)"
                    )
                    print("█" * 100)
                    for group, url, rows, _ in upload_work:
                        stats = new_stats()
                        started = time.time()
                        print(f"\n🧩 {group['name']} | {url}")
                        notes = sync_uploads(page, state, group, url, rows, stats)
                        group_upload_problems[(profile_name, group["name"])] = notes
                        print_group_done(group["name"], stats, started)
                        add_stats(totals, stats)
        except BlockedError as e:
            stopped = str(e)
        except KeyboardInterrupt:
            stopped = "stopped by the user (Ctrl+C)"
        except Exception as e:
            stopped = f"{type(e).__name__}: {e}"
        finally:
            if context is not None:
                try:
                    context.close()
                except Exception:
                    pass

    P.emit("phase", phase="check" if PLAN_ONLY else ("upload" if DO_UPLOADS else "sync"), status="error" if stopped else "done", reason=stopped or "")
    write_plan_csv(plan_rows)
    if not PLAN_ONLY:
        write_results_excel(state)
    if U.MISSING_FROM_WAVZ_ROWS:
        try:
            U.write_missing_from_wavz()
        except Exception as e:
            print(f"⚠️ Could not save the missing-from-Wavz report: {e}")

    print("\n" + "#" * 100)
    if stopped:
        cur = state.data.get("current") or {}
        state.save()
        print(f"⛔ Stopped: {stopped}")
        print(
            f"   Last position: group {cur.get('group')} | step: {cur.get('step')} | session: {cur.get('session')}"
        )
        print("   Run it again to continue from there (finished steps are not repeated).")
    else:
        state.finish()
        print("✅ Finished.")
    if not PLAN_ONLY:
        print(
            f"   created={totals['created']} | edited={totals['edited']} | uploaded={totals['uploaded']} | "
            f"reuploaded={totals['reuploaded']} | failed={totals['failed']}"
        )
        for f in state.data["failed"]:
            print(f"   ❌ {f['group']} | {f['step']} | {f['session']} | {f['reason']}")
    for (profile_name, group_name), notes in group_upload_problems.items():
        problems.extend((profile_name, group_name, n) for n in notes)
    if problems:
        print(f"\n⚠️ Needs your attention ({len(problems)}): these could not be done automatically")
        for profile_name, group_name, n in problems:
            P.TRACK = "nontechnical" if "non" in profile_name.lower() else "technical"
            P.session(group_name, n, "check", "error", n.get("reason", ""))
            d = n.get("date")
            print(
                f"   - {profile_name} | {group_name} | {d.isoformat() if d else ''} | "
                f"{str(n.get('title', ''))[:60]} -> {n.get('reason', '')}"
            )
    else:
        print("✅ Nothing skipped: every session in the CSVs is covered.")
    print(f"   Progress file: {STATE_PATH}")
    print(f"   Total time: {time.time() - run_started:.1f}s")
    print("#" * 100)
    if stopped:
        sys.exit(1)


if __name__ == "__main__":
    main()
